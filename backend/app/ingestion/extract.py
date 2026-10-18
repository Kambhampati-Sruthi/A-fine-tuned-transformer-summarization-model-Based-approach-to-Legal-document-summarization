"""Upload validation + PDF/DOCX text extraction with page-aware offsets."""
from __future__ import annotations
import io, re, unicodedata
from collections import Counter
from dataclasses import dataclass, field
from typing import Optional

import pymupdf  # PyMuPDF
import docx  # python-docx

from .textutil import _is_heading

ALLOWED = {".pdf": b"%PDF-", ".docx": b"PK\x03\x04"}
MIN_CHARS_PER_PAGE = 40  # below this a page is treated as image-only/scanned


class IngestionError(ValueError):
    """User-facing, safe-to-display ingestion problem."""


@dataclass
class ExtractedDoc:
    text: str
    page_spans: list[tuple[Optional[int], int, int]]  # (page, start, end) over `text`
    n_pages: Optional[int]
    warnings: list[str] = field(default_factory=list)

    def page_at(self, offset: int) -> Optional[int]:
        for page, s, e in self.page_spans:
            if s <= offset < e:
                return page
        return self.page_spans[-1][0] if self.page_spans else None


def validate_upload(filename: str, data: bytes, max_mb: int) -> str:
    """Check extension, magic bytes and size. Returns normalised extension."""
    name = (filename or "").lower()
    ext = next((e for e in ALLOWED if name.endswith(e)), None)
    if not ext:
        raise IngestionError("Only .pdf and .docx files are supported.")
    if not data:
        raise IngestionError("The uploaded file is empty.")
    if len(data) > max_mb * 1024 * 1024:
        raise IngestionError(f"File exceeds the {max_mb} MB limit.")
    if not data.startswith(ALLOWED[ext]):
        raise IngestionError("File content does not match its extension.")
    return ext


def _clean_page_text(t: str) -> str:
    t = unicodedata.normalize("NFKC", t)
    t = t.replace("\u00ad", "")                      # soft hyphen
    t = re.sub(r"(\w)-\n(\w)", r"\1\2", t)           # de-hyphenate line breaks
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


LIST_ITEM = re.compile(r"^\(?(?:[a-z]|[ivx]{1,4}|\d{1,2})[).]\s+\S")


def reflow(text: str) -> str:
    """Join PDF soft line-wraps into paragraphs; keep breaks around headings, list items, blank lines."""
    out: list[str] = []
    for line in text.split("\n"):
        if not line.strip():
            if out and out[-1] != "":
                out.append("")
            continue
        if not out or out[-1] == "":
            out.append(line.strip()); continue
        prev = out[-1]
        if _is_heading(line) or _is_heading(prev) or LIST_ITEM.match(line.strip()) or prev.isupper():
            out.append(line.strip())
        else:
            out[-1] = prev + " " + line.strip()
    return "\n".join(out)


def _strip_running_headers(pages: list[str]) -> list[str]:
    """Remove lines (headers/footers/page numbers) repeated on >50% of pages."""
    if len(pages) < 3:
        return pages
    norm = lambda l: re.sub(r"\d+", "#", l.strip().lower())
    counts: Counter = Counter()
    for p in pages:
        counts.update({norm(l) for l in p.splitlines() if l.strip()})
    repeated = {k for k, v in counts.items() if v > len(pages) * 0.5 and len(k) < 80}
    return ["\n".join(l for l in p.splitlines() if norm(l) not in repeated) for p in pages]


def _assemble(pages: list[tuple[Optional[int], str]], warnings, n_pages) -> ExtractedDoc:
    text, spans, pos = "", [], 0
    for page, t in pages:
        if not t:
            continue
        sep = "\n\n" if text else ""
        text += sep
        pos += len(sep)
        spans.append((page, pos, pos + len(t)))
        text += t
        pos += len(t)
    if not text.strip():
        raise IngestionError(
            "No extractable text found. The file may be scanned or image-only; "
            "enable OCR (OCR_ENABLED=true, requires Tesseract) or upload a text-based file."
        )
    return ExtractedDoc(text, spans, n_pages, warnings)


def extract_pdf(data: bytes, max_pages: int = 400, ocr: bool = False) -> ExtractedDoc:
    try:
        pdf = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # malformed / encrypted
        raise IngestionError("The PDF could not be opened (corrupt or malformed).") from exc
    if pdf.needs_pass:
        raise IngestionError("Password-protected PDFs are not supported.")
    if pdf.page_count > max_pages:
        raise IngestionError(f"PDF has {pdf.page_count} pages; the limit is {max_pages}.")
    warnings, raw = [], []
    for i, page in enumerate(pdf, start=1):
        t = page.get_text("text", sort=True)
        if len(t.strip()) < MIN_CHARS_PER_PAGE:
            if ocr:
                try:
                    t = page.get_textpage_ocr(full=True).extractText()
                except Exception:
                    warnings.append(f"Page {i}: OCR failed (is Tesseract installed?).")
            else:
                warnings.append(f"Page {i} has little/no text (possibly scanned); OCR is disabled.")
        raw.append(t)
    cleaned = [reflow(p) for p in _strip_running_headers([_clean_page_text(t) for t in raw])]
    return _assemble([(i, t) for i, t in enumerate(cleaned, 1)], warnings, pdf.page_count)


def extract_docx(data: bytes) -> ExtractedDoc:
    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise IngestionError("The DOCX could not be opened (corrupt or malformed).") from exc
    lines = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    for tbl in d.tables:  # keep tabular content (fee schedules etc.) searchable
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return _assemble([(None, _clean_page_text("\n".join(lines)))], [], None)


def extract(filename: str, data: bytes, max_mb=20, max_pages=400, ocr=False) -> ExtractedDoc:
    ext = validate_upload(filename, data, max_mb)
    return extract_pdf(data, max_pages, ocr) if ext == ".pdf" else extract_docx(data)
