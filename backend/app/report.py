"""Render a Report as Markdown, DOCX or PDF."""
from __future__ import annotations
import io
from .schemas.models import Report

SEV = {"high": "HIGH", "medium": "MEDIUM", "low": "LOW", "info": "INFO"}


def _ev(report: Report, ids, chunks_index):
    return ", ".join(chunks_index.get(i, i) for i in ids)


def to_markdown(r: Report, refs: dict[str, str] | None = None) -> str:
    refs = refs or {}
    L = [f"# Contract Analysis: {r.filename}", f"_{r.disclaimer}_", "",
         "## Executive summary", r.summary.executive_summary, "",
         f"*Summariser: {r.summary.mode} / {r.summary.model} - clause coverage {r.summary.coverage:.0%}*", "",
         "## Key facts"]
    L += [f"- **{f.field}:** {f.value}  _[{_ev(r, f.evidence_ids, refs)}]_" for f in r.key_facts] or ["- none detected"]
    L += ["", "## Section summaries"]
    for s in r.summary.sections:
        L.append(f"- **{s.heading}** - {s.summary}  _[{_ev(r, s.evidence_ids[:3], refs)}]_")
    L += ["", "## Potential issues for review"]
    for k in r.risks:
        L.append(f"- **[{SEV[k.severity]}] {k.title}** ({k.certainty.replace('_', ' ')}) - {k.rationale}"
                 + (f"  _[{_ev(r, k.evidence_ids, refs)}]_" if k.evidence_ids else ""))
    if not r.risks:
        L.append("- none detected by the rule set")
    L += ["", "## Unresolved / unverified"] + ([f"- {u}" for u in r.unresolved] or ["- none"])
    if r.warnings:
        L += ["", "## Warnings"] + [f"- {w}" for w in r.warnings]
    L += ["", f"Review status: **{r.review_status}**" + (f" - {r.reviewer_notes}" if r.reviewer_notes else "")]
    return "\n".join(L)


def to_docx(r: Report, refs: dict[str, str] | None = None) -> bytes:
    import docx
    refs = refs or {}
    d = docx.Document()
    d.add_heading("Contract Analysis Report", 0)
    d.add_paragraph(r.filename).runs[0].bold = True
    d.add_paragraph(r.disclaimer).runs[0].italic = True
    d.add_heading("Executive summary", 1); d.add_paragraph(r.summary.executive_summary)
    d.add_paragraph(f"Summariser: {r.summary.mode} / {r.summary.model}; coverage {r.summary.coverage:.0%}")
    d.add_heading("Key facts", 1)
    for f in r.key_facts:
        d.add_paragraph(f"{f.field}: {f.value}  [{_ev(r, f.evidence_ids, refs)}]", style="List Bullet")
    d.add_heading("Section summaries", 1)
    for s in r.summary.sections:
        d.add_paragraph(f"{s.heading} - {s.summary}  [{_ev(r, s.evidence_ids[:3], refs)}]", style="List Bullet")
    d.add_heading("Potential issues for review", 1)
    t = d.add_table(rows=1, cols=4); t.style = "Light Grid Accent 1"
    for i, h in enumerate(("Severity", "Issue", "Rationale", "Evidence")):
        t.rows[0].cells[i].text = h
    for k in r.risks:
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text = SEV[k.severity], k.title, k.rationale
        row[3].text = _ev(r, k.evidence_ids, refs) or "(absence)"
    d.add_heading("Unresolved / unverified", 1)
    for u in r.unresolved or ["none"]:
        d.add_paragraph(u, style="List Bullet")
    d.add_paragraph(f"Review status: {r.review_status}")
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()


def to_pdf(r: Report, refs: dict[str, str] | None = None) -> bytes:
    from xml.sax.saxutils import escape
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    st = getSampleStyleSheet()
    md = to_markdown(r, refs)
    story = []
    for line in md.splitlines():
        if not line.strip():
            story.append(Spacer(1, 6)); continue
        style, txt = st["BodyText"], line
        if line.startswith("# "): style, txt = st["Title"], line[2:]
        elif line.startswith("## "): style, txt = st["Heading2"], line[3:]
        elif line.startswith("- "): txt = "• " + line[2:]
        txt = escape(txt).replace("**", "").replace("_[", "[").replace("]_", "]").strip("_")
        story.append(Paragraph(txt, style))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title="Contract analysis").build(story)
    return buf.getvalue()
