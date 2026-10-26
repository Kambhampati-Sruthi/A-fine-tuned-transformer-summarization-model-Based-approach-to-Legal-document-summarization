
import io
import re
import sys
import time
import html
import traceback
from pathlib import Path
from transformers import LEDTokenizerFast, LEDForConditionalGeneration
import pandas as pd
import streamlit as st
import torch
from pypdf import PdfReader
from docx import Document
from transformers import LEDTokenizer, LEDForConditionalGeneration

# ------------------------------------------------------------
# Paths and imports
# ------------------------------------------------------------
ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
CHECKPOINT = ROOT / "models" / "led_finetuned" / "checkpoint-147"

if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

# Import the existing retrieval backend without hiding import errors.
try:
    from pipeline.claim_evidence_pipeline import build_evidence_retriever
    from pipeline.contract_analysis import analyze_claims
    BACKEND_IMPORT_ERROR = None
except Exception as exc:
    build_evidence_retriever = None
    analyze_claims = None
    BACKEND_IMPORT_ERROR = f"{type(exc).__name__}: {exc}"

DEVICE = torch.device("cuda" if torch.cuda.is_available() else "cpu")

# ------------------------------------------------------------
# Page configuration and visual design
# ------------------------------------------------------------
st.set_page_config(
    page_title="Legal Document Summarization",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Manrope:wght@400;500;600;700;800&display=swap');

html, body, [class*="css"] {
    font-family: 'DM Sans', sans-serif;
}
.stApp {
    background: #f5f7fb;
}
[data-testid="stSidebar"] {
    background: #111b31;
}
[data-testid="stSidebar"] * {
    color: #eef3ff;
}
[data-testid="stSidebar"] [data-testid="stMarkdownContainer"] p {
    color: #b8c5df;
}
.hero {
    padding: 28px 30px;
    border-radius: 22px;
    color: white;
    background: linear-gradient(120deg, #111b31 0%, #203a67 65%, #355e91 100%);
    margin-bottom: 22px;
}
.hero h1 {
    font-family: 'Manrope', sans-serif;
    font-size: clamp(1.35rem, 2.4vw, 2.1rem);
    line-height: 1.25;
    overflow-wrap: anywhere;
    font-weight: 800;
    letter-spacing: -1px;
    margin: 0 0 8px 0;
    color: white;
}
.hero p {
    color: #d5e2f7;
    margin: 0;
    font-size: 1rem;
}
.eyebrow {
    text-transform: uppercase;
    letter-spacing: 2px;
    color: #9ec5ff;
    font-size: .72rem;
    font-weight: 700;
    margin-bottom: 10px;
}
.panel {
    background: white;
    border: 1px solid #e5eaf3;
    padding: 19px 20px;
    border-radius: 16px;
    margin-bottom: 14px;
}
.metric {
    background: white;
    border: 1px solid #e5eaf3;
    border-radius: 15px;
    padding: 17px;
}
.metric-label {
    font-size: .8rem;
    color: #69758b;
    margin-bottom: 6px;
}
.metric-value {
    font-family: 'Manrope', sans-serif;
    font-size: 1.65rem;
    font-weight: 800;
    color: #16243c;
}
.muted {
    color: #69758b;
    font-size: .88rem;
}
div.stButton > button[kind="primary"] {
    background: #254d83;
    border: 0;
    border-radius: 10px;
    min-height: 44px;
    font-weight: 700;
}
div.stButton > button {
    border-radius: 10px;
}
div[data-testid="stFileUploader"] {
    border-radius: 14px;
}
h2, h3 {
    font-family: 'Manrope', sans-serif;
    color: #17243b;
}
</style>
""", unsafe_allow_html=True)


# ------------------------------------------------------------
# Helpers
# ------------------------------------------------------------
def extract_uploaded_text(uploaded_file):
    name = uploaded_file.name.lower()
    raw = uploaded_file.getvalue()

    if name.endswith(".txt"):
        text = raw.decode("utf-8", errors="replace")
    elif name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(raw))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    else:
        raise ValueError("Supported formats are PDF, DOCX and TXT.")

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    if not text:
        raise ValueError(
            "No text could be extracted. This may be a scanned PDF; "
            "OCR is not enabled in this version."
        )
    return text



def segment_contract_by_heading(text):
    """Create source-linked sections using conservative heading detection."""
    text = str(text)

    # Accept explicit labels, numbered clauses, and short uppercase headings.
    heading_pattern = re.compile(
        r"(?m)^[ \t]*(?P<title>"
        r"(?:Section|Article|Clause)\s+[A-Z0-9IVX.-]+[^\n]{0,100}"
        r"|(?:\d+(?:\.\d+)*[.)]?\s+[A-Z][^\n]{1,90})"
        r"|(?:[A-Z][A-Za-z0-9 /&(),'-]{1,70}:)"
        r"|(?:[A-Z][A-Z0-9 ,&()/.-]{2,70})"
        r")[ \t]*$"
    )

    headings = []

    for match in heading_pattern.finditer(text):
        title = match.group("title").strip()

        # Avoid treating ordinary sentence lines as headings.
        if title.endswith(":"):
            pass
        elif re.match(
            r"^(Section|Article|Clause)\s+",
            title,
            flags=re.IGNORECASE
        ):
            pass
        elif re.match(r"^\d", title):
            pass
        elif title.isupper() and len(title.split()) <= 12:
            pass
        else:
            continue

        headings.append(match)

    if not headings:
        matches = list(re.finditer(r"\S+", text))
        records = []
        max_words, overlap_words = 400, 60
        step = max_words - overlap_words

        for start in range(0, len(matches), step):
            end = min(start + max_words, len(matches))
            start_char = matches[start].start()
            end_char = matches[end - 1].end()

            records.append({
                "evidence_id": f"CHUNK-{len(records) + 1:04d}",
                "section": f"Chunk {len(records) + 1}",
                "title": "",
                "source_start_char": start_char,
                "source_end_char": end_char,
                "text": text[start_char:end_char],
            })

            if end == len(matches):
                break

        return records

    boundaries = []

    if text[:headings[0].start()].strip():
        boundaries.append((0, headings[0].start(), "Introduction"))

    for index, heading in enumerate(headings):
        start = heading.start()
        end = (
            headings[index + 1].start()
            if index + 1 < len(headings)
            else len(text)
        )

        title = heading.group("title").strip().rstrip(":").strip()

        if text[start:end].strip():
            boundaries.append((start, end, title or "Section"))

    records = []

    for start, end, title in boundaries:
        records.append({
            "evidence_id": f"SECTION-{len(records) + 1:04d}",
            "section": title,
            "title": title,
            "source_start_char": start,
            "source_end_char": end,
            "text": text[start:end].strip(),
        })

    return records



def chunk_text(text, max_words=400, overlap_words=60):
    if max_words <= 0 or overlap_words < 0 or overlap_words >= max_words:
        raise ValueError("Overlap must be non-negative and smaller than chunk size.")

    words = text.split()
    if not words:
        return []

    step = max_words - overlap_words
    chunks = []
    for start in range(0, len(words), step):
        part = " ".join(words[start:start + max_words]).strip()
        if part:
            chunks.append(part)
        if start + max_words >= len(words):
            break
    return chunks


def simple_review_flags(source, summary):
    flags = []
    if not summary or len(summary.split()) < 10:
        flags.append("Very short or empty summary")

    source_words = set(re.findall(r"\b\w+\b", source.lower()))
    summary_words = set(re.findall(r"\b\w+\b", summary.lower()))
    meaningful = {
        w for w in source_words
        if len(w) > 3 and w not in {
            "this", "that", "with", "from", "shall", "will",
            "have", "which", "such", "there", "their", "under",
            "agreement", "party", "parties"
        }
    }
    overlap = len(meaningful & summary_words) / len(meaningful) if meaningful else 0
    if overlap < 0.10:
        flags.append("Low source-word overlap")

    # Surface numbers that appear in the source but not in the summary.
    source_numbers = set(re.findall(r"\b\d+(?:\.\d+)?%?\b", source))
    summary_numbers = set(re.findall(r"\b\d+(?:\.\d+)?%?\b", summary))
    missing_numbers = sorted(source_numbers - summary_numbers)
    if missing_numbers:
        flags.append("Some source numbers are absent from the summary")

    return flags or ["No simple flag; correctness is not established"]


@st.cache_resource(show_spinner=False)
def load_legal_model(checkpoint_path):
    #tokenizer = LEDTokenizer.from_pretrained(checkpoint_path)
    tokenizer = LEDTokenizerFast.from_pretrained(checkpoint_path)
    model = LEDForConditionalGeneration.from_pretrained(checkpoint_path)
    model.to(DEVICE)
    model.eval()
    return tokenizer, model


def summarize_legal(text, tokenizer, model, max_new_tokens=110):
    inputs = tokenizer(
        text,
        max_length=512,
        truncation=True,
        return_tensors="pt",
    )
    inputs = {key: value.to(DEVICE) for key, value in inputs.items()}

    global_mask = torch.zeros_like(inputs["attention_mask"])
    global_mask[:, 0] = 1

    with torch.inference_mode():
        output = model.generate(
            input_ids=inputs["input_ids"],
            attention_mask=inputs["attention_mask"],
            global_attention_mask=global_mask,
            max_new_tokens=max_new_tokens,
            num_beams=4,
            no_repeat_ngram_size=3,
        )
    return tokenizer.decode(output[0], skip_special_tokens=True).strip()


def metric_card(label, value, note=""):
    st.markdown(
        f"""<div class="metric">
        <div class="metric-label">{html.escape(str(label))}</div>
        <div class="metric-value">{html.escape(str(value))}</div>
        <div class="muted">{html.escape(str(note))}</div>
        </div>""",
        unsafe_allow_html=True,
    )


def reset_results():
    for key in [
        "result_df", "source_text", "uploaded_name", "elapsed",
        "section_records", "document_chunks", "retriever",
    ]:
        st.session_state.pop(key, None)


# Header
# ------------------------------------------------------------
st.markdown("""
<div class="hero">
  <div class="eyebrow">LEGAL INTELLIGENCE · DOCUMENT WORKSPACE</div>
  <h1>A fine-tuned transformer summarization model-Based Approach to Legal document summarization</h1>
  <p>Summarize legal documents, retrieve relevant clauses, and inspect
  generated statements against source evidence.</p>
</div>
""", unsafe_allow_html=True)

# ------------------------------------------------------------
# Top navigation
# ------------------------------------------------------------
nav_left, nav_center, nav_right = st.columns([1, 3, 1])

with nav_center:
    page = st.radio(
        "Workspace",
        [
            "Document Studio",
            "Evidence Search",
            "Claim Review",
            "Paralegal Review",
        ],
        horizontal=True,
        label_visibility="collapsed",
        key="workspace_navigation",
    )

# Preserve existing processing defaults.
# The controls will be displayed only in Document Studio.
max_words = 400
overlap_words = 60
max_new_tokens = 110


st.info(
    "Research prototype: AI summaries and automated screening can be wrong. "
    "Always inspect the original document and obtain qualified legal review."
)

# ------------------------------------------------------------
# Document Studio
# ------------------------------------------------------------
if page == "Document Studio":
    with st.expander("Processing settings", expanded=False):
        settings_cols = st.columns(3)
        with settings_cols[0]:
            max_words = st.slider("Words per chunk", 200, 500, 400, 50)
        with settings_cols[1]:
            overlap_words = st.slider("Chunk overlap", 0, 100, 60, 10)
        with settings_cols[2]:
            max_new_tokens = st.slider("Summary length", 60, 160, 110, 10)

    st.subheader("Document Studio")
    st.write("Upload a document to extract its text, create source-linked chunks, and generate summaries.")

    uploaded_file = st.file_uploader(
        "Choose a legal document",
        type=["pdf", "docx", "txt"],
        help="PDF, DOCX and TXT are supported. Scanned PDFs may need OCR first.",
    )

    left, right = st.columns([1.25, 1])
    with left:
        st.markdown("#### What this workspace does")
        st.markdown("""
        - Extracts document text
        - Splits the document into overlapping chunks
        - Generates summaries using your fine-tuned LED model
        - Displays the original source beside each summary
        - Provides downloadable review records
        """)
    with right:
        st.markdown("#### Model status")
        if CHECKPOINT.exists():
            st.success("Fine-tuned LED checkpoint found")
        else:
            st.error(f"Checkpoint missing: {CHECKPOINT}")

        if BACKEND_IMPORT_ERROR:
            st.warning("The evidence-search or claim-review backend could not be imported.")
        else:
            st.success("Retrieval and claim-review modules imported")

    if uploaded_file is not None:
        st.markdown("---")
        st.write(f"**Selected file:** `{uploaded_file.name}`")
        st.write(f"**File size:** {uploaded_file.size / 1024:.1f} KB")

        if st.button("Extract, summarize and index document", type="primary", use_container_width=True):
            if overlap_words >= max_words:
                st.error("Chunk overlap must be smaller than the chunk size.")
            elif not CHECKPOINT.exists():
                st.error(f"Model checkpoint not found at `{CHECKPOINT}`.")
            else:
                reset_results()
                try:
                    with st.spinner("Extracting document text..."):
                        full_text = extract_uploaded_text(uploaded_file)
                        chunks = chunk_text(full_text, max_words, overlap_words)

                    if not chunks:
                        st.error("No usable text chunks were created.")
                        st.stop()

                    progress = st.progress(0, text="Loading the fine-tuned LED model...")
                    tokenizer, model = load_legal_model(str(CHECKPOINT))

                    records = []
                    start_time = time.time()
                    for i, chunk in enumerate(chunks):
                        try:
                            summary = summarize_legal(
                                chunk, tokenizer, model, max_new_tokens
                            )
                            flags = simple_review_flags(chunk, summary)
                            records.append({
                                "chunk_index": i + 1,
                                "source_word_count": len(chunk.split()),
                                "generated_summary": summary,
                                "review_flags": "; ".join(flags),
                                "verification_status": "Human review required",
                                "source_text": chunk,
                            })
                        except Exception as exc:
                            records.append({
                                "chunk_index": i + 1,
                                "source_word_count": len(chunk.split()),
                                "generated_summary": "",
                                "review_flags": f"Generation failed: {type(exc).__name__}: {exc}",
                                "verification_status": "Generation failed",
                                "source_text": chunk,
                            })
                        progress.progress(
                            (i + 1) / len(chunks),
                            text=f"Summarizing chunk {i + 1} of {len(chunks)}...",
                        )

                    elapsed = time.time() - start_time
                    result_df = pd.DataFrame(records)

                    # Build a heading-aware source index.
                    # Summarization chunks and the LED model remain unchanged.
                    section_records = segment_contract_by_heading(full_text)

                    st.session_state["result_df"] = result_df
                    st.session_state["source_text"] = full_text
                    st.session_state["uploaded_name"] = uploaded_file.name
                    st.session_state["elapsed"] = elapsed
                    st.session_state["section_records"] = section_records
                    st.session_state["document_chunks"] = chunks
                    st.session_state["retriever"] = None

                    progress.empty()
                    st.success("Document processing finished.")
                except Exception as exc:
                    st.error(f"Processing failed: {type(exc).__name__}: {exc}")
                    with st.expander("Technical details"):
                        st.code(traceback.format_exc())

    if "result_df" in st.session_state:
        df = st.session_state["result_df"].copy()
        successful = int(
            df["generated_summary"].fillna("").str.strip().ne("").sum()
        )
        failed = len(df) - successful
        word_count = len(st.session_state["source_text"].split())

        # Build a readable overview from the summaries already generated.
        summary_parts = [
            str(value).strip()
            for value in df["generated_summary"].fillna("")
            if str(value).strip()
        ]
        document_summary = "\n\n".join(summary_parts)

        st.markdown("---")
        st.markdown(
            """
            <div style="
                padding: 24px 26px;
                border-radius: 18px;
                background: linear-gradient(135deg, #101b33 0%, #233b60 100%);
                color: white;
                margin-bottom: 18px;">
                <div style="font-size: 12px; letter-spacing: 2px;
                            text-transform: uppercase; opacity: .8;">
                    LEGAL DOCUMENT INTELLIGENCE
                </div>
                <div style="font-size: 29px; font-weight: 700;
                            margin-top: 8px;">
                    Executive Summary
                </div>
                <div style="font-size: 14px; opacity: .85; margin-top: 6px;">
                    Generated summaries for the processed document.
                </div>
            </div>
            """,
            unsafe_allow_html=True,
        )

        st.caption(f"Document: {st.session_state['uploaded_name']}")

        if document_summary:
            st.markdown(document_summary)
            st.download_button(
                "Download summary (.txt)",
                data=(
                    f"Document: {st.session_state['uploaded_name']}\n\n"
                    "EXECUTIVE SUMMARY\n"
                    "Generated from the document's processed text chunks.\n\n"
                    + document_summary
                ).encode("utf-8"),
                file_name="legal_document_summary.txt",
                mime="text/plain",
                type="primary",
            )
        else:
            st.warning(
                "A summary could not be generated. "
                "Please retry processing the document."
            )

        if failed:
            st.warning(
                f"{failed} section summary generation(s) failed. "
                "The overview may be incomplete."
            )

        st.markdown("### Document overview")
        metrics = st.columns(3)
        with metrics[0]:
            metric_card("Source words", f"{word_count:,}")
        with metrics[1]:
            metric_card("Chunks processed", len(df))
        with metrics[2]:
            metric_card("Summaries generated", successful)

        st.markdown("---")

        with st.expander("View processing details"):
            st.caption(
                "Expand a chunk to compare its generated summary "
                "with its original source passage."
            )

            search = st.text_input(
                "Search chunks",
                placeholder="Try payment, termination or liability...",
                key="section_search",
            )

            filtered = df
            if search.strip():
                mask = (
                    df["generated_summary"].fillna("").str.contains(
                        search, case=False, regex=False
                    )
                    | df["source_text"].fillna("").str.contains(
                        search, case=False, regex=False
                    )
                )
                filtered = df[mask]

            st.caption(f"Showing {len(filtered)} of {len(df)} chunks")

            for _, row in filtered.iterrows():
                with st.expander(
                    f"Chunk {int(row['chunk_index'])}",
                    expanded=False,
                ):
                    st.markdown("**Generated chunk summary**")
                    summary = str(row["generated_summary"] or "").strip()
                    if summary:
                        st.write(summary)
                    else:
                        st.info("No summary was generated for this chunk.")

                    flags = str(row.get("review_flags", "") or "").strip()
                    if flags:
                        st.markdown("**Automated quality flags**")
                        st.caption(flags)

                    st.markdown("**Original source text**")
                    st.text_area(
                        "Source passage",
                        value=str(row["source_text"] or ""),
                        height=200,
                        key=f"source_{row['chunk_index']}",
                        label_visibility="collapsed",
                    )

        with st.expander("Export detailed analysis"):
            csv_data = df.to_csv(index=False).encode("utf-8")
            st.download_button(
                "Download source-linked records (.csv)",
                data=csv_data,
                file_name="source_linked_summary_review.csv",
                mime="text/csv",
            )

            report_lines = [
                f"Document: {st.session_state['uploaded_name']}",
                "CHUNK SUMMARIES",
                "=" * 60,
                "",
            ]
            for _, row in df.iterrows():
                report_lines.extend([
                    f"Chunk {row['chunk_index']}",
                    "",
                    str(row["generated_summary"] or "No summary available."),
                    "",
                    "Original source:",
                    str(row["source_text"] or ""),
                    "",
                    "-" * 60,
                    "",
                ])

            st.download_button(
                "Download detailed report (.txt)",
                data="\n".join(report_lines).encode("utf-8"),
                file_name="legal_document_detailed_report.txt",
                mime="text/plain",
            )



# ------------------------------------------------------------
# Paralegal Review
# ------------------------------------------------------------
elif page == "Paralegal Review":
    st.subheader("Paralegal Review")
    st.write(
        "Review possible topic-coverage gaps in generated summaries "
        "against passages from the original document."
    )
    st.warning(
        "Human review required. Automated output is not a legal opinion. "
        "These flags do not establish legal truth or confirm an omission."
    )

    if "result_df" not in st.session_state or "source_text" not in st.session_state:
        st.info("First upload and process a document in Document Studio.")
    else:
        source_text = str(st.session_state.get("source_text", ""))
        result_df = st.session_state["result_df"]

        summary_parts = [
            str(value).strip()
            for value in result_df["generated_summary"].fillna("")
            if str(value).strip()
        ]
        document_summary = "\n\n".join(summary_parts)

        st.markdown("### Generated summary")
        if document_summary:
            st.write(document_summary)
        else:
            st.info("No generated summary is available.")

        topic_terms = {
            "Arbitration opt-out": [
                "arbitration opt out", "arbitration opt-out",
                "opt out", "opt-out"
            ],
            "Class-action waiver": [
                "class action", "class or representative proceeding"
            ],
            "Governing law and jurisdiction": [
                "governing law", "jurisdiction", "venue"
            ],
            "Termination or suspension": [
                "termination", "terminate", "suspend"
            ],
            "Fees and payments": [
                "filing fee", "administrative fee",
                "arbitrator fees", "payment", "refund", "fees"
            ],
            "Liability and indemnification": [
                "limitation of liability", "indemnification",
                "liability", "damages"
            ],
        }

        review_questions = {
            "Arbitration opt-out":
                "Are the opt-out method, deadline, and consequences preserved?",
            "Class-action waiver":
                "Are the waiver and any related conditions or exceptions preserved?",
            "Governing law and jurisdiction":
                "Are the applicable law and jurisdiction details preserved?",
            "Termination or suspension":
                "Are the relevant termination conditions and consequences preserved?",
            "Fees and payments":
                "Are fee allocations, thresholds, and exceptions preserved?",
            "Liability and indemnification":
                "Are relevant liability limits, exceptions, and obligations preserved?",
        }

        st.markdown("### Coverage review")
        st.caption(
            "Flags use exact keyword matching in the summary and source. "
            "Different wording may be missed, and keyword matches do not "
            "prove that a topic is adequately covered."
        )

        for topic, terms in topic_terms.items():
            matching_positions = []
            source_lower = source_text.lower()

            for term in terms:
                start = 0
                while True:
                    pos = source_lower.find(term, start)
                    if pos < 0:
                        break
                    matching_positions.append((pos, term))
                    start = pos + len(term)

            if not matching_positions:
                continue

            summary_lower = document_summary.lower()
            summary_mentions = any(term in summary_lower for term in terms)

            status = (
                "Topic term found in summary — review content"
                if summary_mentions
                else "Potential coverage gap — human review required"
            )

            with st.expander(f"{topic} — {status}", expanded=False):
                st.write("**Review question:**", review_questions[topic])
                st.write(
                    "**Summary keyword check:**",
                    "At least one configured topic term found."
                    if summary_mentions
                    else "No configured topic term found in the summary."
                )

                st.markdown("**Source evidence**")
                seen_ranges = []
                shown = 0

                for pos, term in sorted(matching_positions):
                    start = max(0, pos - 220)
                    end = min(len(source_text), pos + len(term) + 220)

                    if any(start < old_end and end > old_start
                           for old_start, old_end in seen_ranges):
                        continue

                    excerpt = " ".join(source_text[start:end].split())
                    st.caption(
                        f"Matched term: {term} | "
                        f"Source characters {start}–{end}"
                    )
                    st.text(excerpt)
                    seen_ranges.append((start, end))
                    shown += 1

                    if shown >= 3:
                        break

        st.markdown("---")
        st.caption(
            "Evidence excerpts are text matches from the uploaded document. "
            "Review the full surrounding clause before drawing conclusions."
        )


# Evidence Search / RAG
# ------------------------------------------------------------
elif page == "Evidence Search":
    st.subheader("Evidence Search")
    st.write("Search the uploaded document for relevant passages. Results are source excerpts, not legal conclusions.")

    if "section_records" not in st.session_state:
        st.info("First upload and process a document in Document Studio.")
    elif BACKEND_IMPORT_ERROR:
        st.error("The retrieval backend did not import successfully.")
        st.code(BACKEND_IMPORT_ERROR)
    else:
        query = st.text_input(
            "What would you like to find?",
            placeholder="e.g. What law governs this agreement?",
        )
        top_k = st.slider("Number of evidence passages", 1, 10, 5)

        if st.button("Search document", type="primary"):
            if not query.strip():
                st.warning("Enter a search query first.")
            else:
                try:
                    with st.spinner("Searching document evidence..."):
                        retriever = build_evidence_retriever(
                            st.session_state["section_records"]
                        )
                        hits = retriever.search(query.strip(), top_k=top_k)

                    st.session_state["last_search_query"] = query.strip()
                    st.session_state["last_search_hits"] = hits

                    if not hits:
                        st.warning("No matching passages were returned.")
                    else:
                        st.success(f"Retrieved {len(hits)} candidate passage(s).")
                except Exception as exc:
                    st.error(f"Retrieval failed: {type(exc).__name__}: {exc}")
                    with st.expander("Technical details"):
                        st.code(traceback.format_exc())

        for idx, hit in enumerate(st.session_state.get("last_search_hits", []), 1):
            with st.container(border=True):
                st.markdown(f"**Evidence result {idx} · {hit.get('section', 'Source passage')}**")
                if hit.get("title"):
                    st.caption(str(hit["title"]))
                score = hit.get("retrieval_score")
                if score is not None:
                    st.caption(f"Retrieval score: {float(score):.3f} · Not a probability")
                st.write(hit.get("text", ""))
                st.caption(f"Evidence ID: {hit.get('evidence_id', 'Not provided')}")


# ------------------------------------------------------------
# Claim Review
# ------------------------------------------------------------
elif page == "Claim Review":
    st.subheader("Claim Review")
    st.write(
        "Compare a user-entered claim with retrieved source passages using the existing "
        "screening pipeline. This is a preliminary lexical/numeric screen, not legal verification."
    )

    if "section_records" not in st.session_state:
        st.info("First upload and process a document in Document Studio.")
    elif BACKEND_IMPORT_ERROR or analyze_claims is None:
        st.error("Claim-review pipeline is unavailable.")
        st.code(BACKEND_IMPORT_ERROR or "analyze_claims could not be imported.")
    else:
        claim = st.text_area(
            "Claim to examine",
            placeholder="e.g. The agreement is governed by California law.",
            height=100,
        )
        top_k = st.slider("Evidence candidates for claim review", 1, 10, 5, key="claim_top_k")

        if st.button("Retrieve evidence and screen claim", type="primary"):
            if not claim.strip():
                st.warning("Enter a claim first.")
            else:
                try:
                    with st.spinner("Retrieving evidence and running preliminary checks..."):
                        analysis = analyze_claims(
                            st.session_state["section_records"],
                            [claim.strip()],
                            top_k=top_k,
                        )
                    st.session_state["last_claim_analysis"] = analysis
                except Exception as exc:
                    st.error(f"Claim review failed: {type(exc).__name__}: {exc}")
                    with st.expander("Technical details"):
                        st.code(traceback.format_exc())

        analysis = st.session_state.get("last_claim_analysis")
        if analysis:
            for item in analysis.get("results", []):
                status = item.get("status", "UNRESOLVED")
                if status == "SUPPORTED":
                    st.success("Preliminary screen: SUPPORTED")
                elif status == "CONTRADICTED":
                    st.error("Preliminary screen: CONTRADICTED")
                else:
                    st.warning("Preliminary screen: UNRESOLVED")

                st.caption(
                    "These labels reflect heuristic screening only. "
                    "They do not establish legal truth."
                )
                st.write(item.get("reason", "No explanation was returned."))
                st.markdown("**Retrieved evidence**")

                candidates = item.get("retrieval_candidates", [])
                if not candidates:
                    st.info("No retrieval candidates were returned.")
                for i, evidence in enumerate(candidates, 1):
                    with st.expander(
                        f"Evidence {i} · {evidence.get('section', 'Passage')}"
                    ):
                        st.write(evidence.get("text", ""))
                        st.caption(
                            f"Evidence ID: {evidence.get('evidence_id', 'Not provided')}"
                        )

                st.warning(
                    "Human review required. Automated output is not a legal opinion."
                )


# ------------------------------------------------------------
